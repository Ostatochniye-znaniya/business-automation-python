"""Printable forms populated from persisted domain data."""

import csv
import io
from copy import deepcopy
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.table import _Row
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

TEMPLATE = Path(__file__).with_name("templates") / "group_report.docx"


def _cell_text(cell, text: str) -> None:
    """Replace text only, retaining the reference template's paragraph/run properties."""
    nodes = list(cell._tc.iter(qn("w:t")))
    if nodes:
        nodes[0].text = text
        nodes[0].set(qn("xml:space"), "preserve")
        for node in nodes[1:]:
            node.text = ""
    else:
        cell.paragraphs[0].add_run(text)


def report_details(context: dict) -> list[str]:
    return [
        f"Учебный год: {context.get('academic_year', '')}",
        f"Код и наименование направления (специальности): {context.get('program_code', '')}",
        f"Наименование образовательной программы: {context.get('program_name', '')}",
        f"Наименование дисциплины: {context.get('discipline', '')}",
        "ФИО педагогического работника, проводившего занятия в предыдущем семестре: "
        + (context.get("previous_teacher") or ""),
        f"Учебная группа № {context.get('group', '')}",
    ]


def result_values(row: dict) -> list[str]:
    value = row.get("testing_result")
    return [
        row["student"],
        str(row["intermediate_result"]) if row.get("intermediate_result") is not None else "",
        {"ABSENT": "Не явился", "NOT_ADMITTED": "Не допущен"}.get(
            row.get("participation_status"), str(value) if value is not None else ""
        ),
    ]


def report_summaries(rows: list[dict], template: bool) -> list[str]:
    counts = [
        len(rows),
        *[
            sum(row.get("participation_status") == code for row in rows)
            for code in ["PASSED", "NOT_ADMITTED", "ABSENT"]
        ],
    ]
    return [
        f"{label}: {count if not template or index == 0 else '__________'} (чел.)"
        for index, (label, count) in enumerate(
            zip(
                [
                    "Всего в группе",
                    "Прошли тестирование",
                    "Не допущены к тестированию",
                    "Не явилось на тестирование",
                ],
                counts,
                strict=True,
            )
        )
    ]


def report_docx(context: dict, rows: list[dict], template: bool = False) -> bytes:
    document = Document(TEMPLATE)
    table = document.tables[0]
    for index, text in enumerate(report_details(context)):
        _cell_text(table.rows[index].cells[-1], text)
    prototype = deepcopy(table.rows[7]._tr)
    for row in list(table.rows)[7:-4]:
        table._tbl.remove(row._tr)
    summary = list(table.rows)[-4:]
    for index, values in enumerate(rows, 1):
        element = deepcopy(prototype)
        summary[0]._tr.addprevious(element)
        row = _Row(element, table)
        contents = (
            [str(index), values["student"], "", ""]
            if template
            else [str(index), *result_values(values)]
        )
        for cell, value in zip(row.cells, contents, strict=True):
            _cell_text(cell, value)
    for row, text in zip(summary, report_summaries(rows, template), strict=True):
        _cell_text(row.cells[0], text)
    signatures = document.tables[1]
    for index, label, key, name in [
        (0, "Заведующий кафедрой", "department", "chair_head"),
        (2, "Декан факультета/Директор института", "faculty", "faculty_head"),
    ]:
        _cell_text(signatures.rows[index].cells[0], f"{label} {context.get(key, '')}")
        if context.get(name):
            _cell_text(signatures.rows[index].cells[-1], context[name])
    buffer = io.BytesIO()
    document.save(buffer)
    # python-docx normalizes XML declarations in unrelated parts on save.
    # Only the populated body is needed; retain every other original package part verbatim.
    output = io.BytesIO()
    with ZipFile(TEMPLATE) as source, ZipFile(buffer) as populated, ZipFile(output, "w") as result:
        for item in source.infolist():
            result.writestr(
                item,
                populated.read(item.filename)
                if item.filename == "word/document.xml"
                else source.read(item.filename),
            )
    return output.getvalue()


def pdf_font() -> str:
    for path in [
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf"),
        Path("C:/Windows/Fonts/times.ttf"),
    ]:
        if path.is_file():
            if "DocumentFont" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("DocumentFont", str(path)))
            return "DocumentFont"
    raise RuntimeError("Для печатных форм требуется шрифт DejaVu Serif или Times New Roman")


def report_flowables(context: dict, rows: list[dict], template: bool = False) -> list:
    style = ParagraphStyle("Report", fontName=pdf_font(), fontSize=9, leading=11)

    def paragraph(text):
        return Paragraph(escape(str(text)), style)

    data = [[paragraph("№"), paragraph(text), "", ""] for text in report_details(context)]
    data.append(
        [
            paragraph(text)
            for text in [
                "№",
                "ФИО обучающегося",
                "Результаты промежуточной аттестации",
                "Результаты тестирования",
            ]
        ]
    )
    for index, row in enumerate(rows, 1):
        values = [row["student"], "", ""] if template else result_values(row)
        number = ["1", "2", "3", "…", "n"][index - 1] if context.get("blank") else str(index)
        data.append([paragraph(text) for text in [number, *values]])
    summary_start = len(data)
    summaries = report_summaries(rows, template)
    if context.get("blank"):
        summaries[0] = "Всего в группе: __________ (чел.)"
    data.extend([[paragraph(text), "", "", ""] for text in summaries])
    table = Table(data, colWidths=[25, 215, 115, 115], repeatRows=7)
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("SPAN", (0, 0), (0, 5)),
                *[("SPAN", (1, index), (3, index)) for index in range(6)],
                *[("SPAN", (0, index), (3, index)) for index in range(summary_start, len(data))],
            ]
        )
    )
    return [
        paragraph("Отчет о проверке остаточных знаний"),
        Spacer(1, 10),
        table,
        Spacer(1, 12),
        paragraph(
            f"Заведующий кафедрой {context.get('department', '')} __________ / {context.get('chair_head', 'ФИО')}"
        ),
        Spacer(1, 10),
        paragraph(
            f"Декан факультета/Директор института {context.get('faculty', '')} __________ / {context.get('faculty_head', 'ФИО')}"
        ),
    ]


def report_pdf(context: dict, rows: list[dict], template: bool = False) -> bytes:
    buffer = io.BytesIO()
    SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=55, rightMargin=55, topMargin=40, bottomMargin=40
    ).build(report_flowables(context, rows, template))
    return buffer.getvalue()


def schedule_pdf(title: str, rows: list[list[str]], signatory: str) -> bytes:
    font_paths = [
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
    ]
    font_path = next((path for path in font_paths if path.is_file()), None)
    if font_path is None:
        raise RuntimeError("Для PDF установите шрифт DejaVu Sans")
    if "ScheduleFont" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("ScheduleFont", str(font_path)))
    style = ParagraphStyle("Schedule", fontName="ScheduleFont", fontSize=8, leading=11)
    header = [
        "№",
        "Группа",
        "Профиль подготовки",
        "Дисциплина",
        "Кафедра",
        "ФИО ППС",
        "Дата / время (Москва)",
        "Место",
    ]
    data = [[Paragraph(escape(str(cell)), style) for cell in row] for row in [header, *rows]]
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4), rightMargin=20, leftMargin=20, topMargin=25, bottomMargin=25
    )
    table = Table(data, repeatRows=1, colWidths=[25, 55, 115, 135, 100, 110, 90, 65])
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    doc.build(
        [
            Paragraph(escape(title), style),
            Spacer(1, 12),
            table,
            Spacer(1, 24),
            Paragraph(escape(signatory), style),
        ]
    )
    return buffer.getvalue()


def csv_bytes(headers: list[str], rows: list[list]) -> bytes:
    def safe(value):
        text = str(value if value is not None else "")
        return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text

    stream = io.StringIO()
    writer = csv.writer(stream, delimiter=";")
    writer.writerow(headers)
    writer.writerows([[safe(value) for value in row] for row in rows])
    return stream.getvalue().encode("utf-8-sig")
